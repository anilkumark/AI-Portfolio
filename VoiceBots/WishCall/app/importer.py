"""Excel / CSV import: read, check, preview, then save."""
import csv
import io
import json
import secrets

from openpyxl import Workbook, load_workbook

from . import db, logic

COLUMNS = ["name", "phone", "language", "group", "occasion", "date", "detail", "repeats", "call_time", "message", "notes"]
REQUIRED = ["name", "phone", "occasion", "date"]
HEADER_ALIASES = {
    "full name": "name", "person": "name", "mobile": "phone", "mobile number": "phone", "phone number": "phone",
    "number": "phone", "lang": "language", "type": "occasion", "occasion type": "occasion", "event": "occasion",
    "group name": "group", "category": "group", "time": "call_time", "call time": "call_time",
    "repeat": "repeats", "yearly": "repeats", "custom message": "message", "note": "notes",
}
MAX_ROWS = 200_000


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(COLUMNS)
    ws.append(["Ravi Kumar", "9848024521", "te", "Friends", "birthday", "14-03-1990", "", "yearly", "09:30", "", ""])
    ws.append(["Priya Sharma", "9900011180", "hi", "Work", "anniversary", "02-12", "", "yearly", "", "", ""])
    ws.append(["Suresh Reddy", "9876500094", "te", "Community", "meeting", "29-10-2026", "6:00 PM at Banjara Hills",
               "once", "16:00", "", ""])
    for col, width in zip("ABCDEFGHIJK", [20, 14, 10, 12, 18, 14, 26, 9, 10, 30, 20]):
        ws.column_dimensions[col].width = width
    info = wb.create_sheet("How to fill")
    rows = [
        ["Column", "What to put"],
        ["name", "Full name (required)"],
        ["phone", "10-digit mobile, e.g. 9848024521 (required). +91 is added automatically."],
        ["language", "en, hi or te (English, Hindi, Telugu). Default: en"],
        ["group", "Family, Friends, Work, Community or Other. Default: Friends"],
        ["occasion", "birthday, anniversary, marriage, condolence, death anniversary, meeting, achievement, festival, custom"],
        ["date", "dd-mm-yyyy, or dd-mm if the year is unknown (yearly occasions only)"],
        ["detail", "Condolence / death anniversary: name of the person. Meeting: time and place. "
                   "Achievement: what they achieved. Festival: festival name."],
        ["repeats", "yearly or once. Leave empty to use the usual for that occasion."],
        ["call_time", "Optional, e.g. 09:30 or 4 PM. Empty = your daily start time."],
        ["message", "Optional. Your own words instead of the template."],
        ["notes", "Optional notes about the person."],
        ["", "One row per occasion. Same phone number on several rows = same person with several occasions."],
    ]
    for r in rows:
        info.append(r)
    info.column_dimensions["A"].width = 12
    info.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_rows(filename: str, data: bytes):
    name = filename.lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    elif name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raise ValueError("Please upload an .xlsx or .csv file.")
    rows = [r for r in rows if r and any(v not in (None, "") for v in r)]
    if not rows:
        raise ValueError("The file is empty.")
    header = []
    for h in rows[0]:
        k = logic.norm_key(h)
        header.append(HEADER_ALIASES.get(k, k.replace(" ", "_")))
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise ValueError("Missing column(s): " + ", ".join(missing) + ". Download the template to see the layout.")
    if len(rows) - 1 > MAX_ROWS:
        raise ValueError(f"Too many rows ({len(rows) - 1:,}). Split the file into parts of {MAX_ROWS:,} rows.")
    out = []
    for i, r in enumerate(rows[1:], start=2):
        rec = {h: (r[j] if j < len(r) else None) for j, h in enumerate(header)}
        rec["_row"] = i
        out.append(rec)
    return out


def _clean(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return s[:-2] if s.endswith(".0") and s[:-2].isdigit() else s


def check(filename: str, data: bytes) -> dict:
    """Validate every row. Returns a summary and stores the good rows for the confirm step."""
    rows = _read_rows(filename, data)
    ready, bad = [], []
    seen = set()
    dup_count = 0
    with db.db() as c:
        existing = {(r[0], r[1], r[2], r[3], r[4] if r[4] is not None else -1) for r in c.execute(
            "SELECT p.phone, o.type, o.month, o.day, o.year FROM occasions o JOIN people p ON p.id = o.person_id")}
    for rec in rows:
        problems = []
        name = _clean(rec.get("name"))
        phone = logic.normalize_phone(rec.get("phone"))
        occ = logic.OCCASION_ALIASES.get(logic.norm_key(rec.get("occasion")))
        lang = logic.LANG_ALIASES.get(logic.norm_key(rec.get("language")) or "en")
        grp_raw = _clean(rec.get("group")).title()
        grp = grp_raw if grp_raw in logic.GROUPS else ("Friends" if not grp_raw else "Other")
        parsed = logic.parse_date(rec.get("date"))
        if not name:
            problems.append("name missing")
        if not phone:
            problems.append("phone not valid")
        if not occ:
            problems.append("occasion not recognised")
        if lang is None:
            problems.append("language must be en, hi or te")
        if not parsed:
            problems.append("date not readable")
        rep = logic.norm_key(rec.get("repeats"))
        recurring = None
        if rep in ("yearly", "every year", "yes", "y", "annual", "1", "true"):
            recurring = True
        elif rep in ("once", "only once", "no", "n", "0", "false"):
            recurring = False
        if occ and recurring is None:
            recurring = logic.OCCASIONS[occ]["yearly"]
        month = day = year = None
        if parsed:
            month, day, year = parsed
            if not logic.valid_day(month, day, year):
                problems.append("date does not exist")
            elif not recurring and year is None:
                problems.append("one-time occasions need the year")
        call_time = None
        if _clean(rec.get("call_time")):
            call_time = logic.parse_time(rec.get("call_time"))
            if not call_time:
                problems.append("call time not readable")
        if problems:
            bad.append({"row": rec["_row"], "name": name, "phone": _clean(rec.get("phone")),
                        "occasion": _clean(rec.get("occasion")), "date": _clean(rec.get("date")),
                        "problem": "; ".join(problems)})
            continue
        key = (phone, occ, month, day, year if year is not None else -1)
        if key in existing or key in seen:
            dup_count += 1
            continue
        seen.add(key)
        ready.append({"name": name, "phone": phone, "language": lang, "grp": grp, "type": occ,
                      "month": month, "day": day, "year": year, "recurring": 1 if recurring else 0,
                      "call_time": call_time, "detail": _clean(rec.get("detail")),
                      "custom_message": _clean(rec.get("message")), "notes": _clean(rec.get("notes")),
                      "ai": 1 if logic.OCCASIONS[occ]["ai"] else 0})
    token = secrets.token_hex(12)
    (db.IMPORTS / f"{token}.json").write_text(json.dumps({"filename": filename, "ready": ready, "bad": bad}),
                                              encoding="utf-8")
    return {"token": token, "filename": filename, "total": len(rows), "ready": len(ready),
            "duplicates": dup_count, "bad": bad, "sample": ready[:15]}


def load(token: str) -> dict | None:
    if not token.isalnum():
        return None
    f = db.IMPORTS / f"{token}.json"
    if not f.exists():
        return None
    return json.loads(f.read_text(encoding="utf-8"))


def errors_csv(token: str) -> str:
    data = load(token) or {"bad": []}
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["row", "name", "phone", "occasion", "date", "problem"])
    for b in data["bad"]:
        w.writerow([b["row"], b["name"], b["phone"], b["occasion"], b["date"], b["problem"]])
    return buf.getvalue()


def commit(token: str) -> tuple[int, int]:
    """Save the checked rows. Returns (new people, new occasions)."""
    data = load(token)
    if not data:
        raise ValueError("This import has expired. Please upload the file again.")
    rows = data["ready"]
    new_people = 0
    with db.db() as c:
        ids = {r["phone"]: r["id"] for r in c.execute("SELECT id, phone FROM people")}
        for r in rows:
            if r["phone"] not in ids:
                cur = c.execute("INSERT INTO people(name, phone, language, grp, notes) VALUES(?,?,?,?,?)",
                                (r["name"], r["phone"], r["language"], r["grp"], r["notes"]))
                ids[r["phone"]] = cur.lastrowid
                new_people += 1
        c.executemany(
            "INSERT INTO occasions(person_id, type, month, day, year, recurring, call_time, detail, custom_message, "
            "ai_replies) VALUES(?,?,?,?,?,?,?,?,?,?)",
            [(ids[r["phone"]], r["type"], r["month"], r["day"], r["year"], r["recurring"], r["call_time"],
              r["detail"], r["custom_message"], r["ai"]) for r in rows])
    (db.IMPORTS / f"{token}.json").unlink(missing_ok=True)
    return new_people, len(rows)
